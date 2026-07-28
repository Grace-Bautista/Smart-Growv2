from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


OUT = Path("analysis_outputs/Smart_Grow_Recommendations_Corrections.docx")
TABLE_DIR = Path("analysis_outputs/table_images")
TABLE_COUNT = 0


def font(size=22, bold=False):
    base = "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"
    return ImageFont.truetype(base, size)


def wrap(draw, text, fnt, width):
    lines = []
    for para in str(text).split("\n"):
        words = para.split()
        if not words:
            lines.append("")
            continue
        line = words[0]
        for word in words[1:]:
            trial = f"{line} {word}"
            if draw.textbbox((0, 0), trial, font=fnt)[2] <= width:
                line = trial
            else:
                lines.append(line)
                line = word
        lines.append(line)
    return lines


def table(doc, headers, rows, widths=None):
    global TABLE_COUNT
    TABLE_COUNT += 1
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    rows = [headers] + rows
    cols = len(headers)
    if widths is None:
        widths = [1] * cols
    total = sum(widths)
    img_w = 1500
    margins = 20
    col_ws = [int((img_w - margins * 2) * w / total) for w in widths]
    col_ws[-1] = img_w - margins * 2 - sum(col_ws[:-1])
    f_body = font(20)
    f_head = font(21, True)
    line_gap = 7
    pad = 12
    scratch = Image.new("RGB", (img_w, 100), "white")
    draw = ImageDraw.Draw(scratch)
    row_heights = []
    wrapped = []
    for ri, row in enumerate(rows):
        row_lines = []
        max_lines = 1
        for ci, cell in enumerate(row):
            f = f_head if ri == 0 else f_body
            lines = wrap(draw, cell, f, col_ws[ci] - pad * 2)
            row_lines.append(lines)
            max_lines = max(max_lines, len(lines))
        wrapped.append(row_lines)
        row_heights.append(max(54, pad * 2 + max_lines * (24 + line_gap)))
    img_h = margins * 2 + sum(row_heights)
    im = Image.new("RGB", (img_w, img_h), "white")
    draw = ImageDraw.Draw(im)
    y = margins
    for ri, row_lines in enumerate(wrapped):
        x = margins
        fill = (217, 234, 247) if ri == 0 else (255, 255, 255)
        for ci, lines in enumerate(row_lines):
            draw.rectangle([x, y, x + col_ws[ci], y + row_heights[ri]], fill=fill, outline=(90, 90, 90), width=2)
            ty = y + pad
            f = f_head if ri == 0 else f_body
            for line in lines:
                draw.text((x + pad, ty), line, fill=(20, 20, 20), font=f)
                ty += 24 + line_gap
            x += col_ws[ci]
        y += row_heights[ri]
    out = TABLE_DIR / f"table_{TABLE_COUNT:02d}.png"
    im.save(out)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(out), width=Inches(7.0))
    doc.add_paragraph()
    return out


def recommendation_block(doc, number, location, issue, action, replacement):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run(f"{number}. {location}")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121)
    for label, value in [
        ("Issue", issue),
        ("Action", action),
        ("Exact replacement/correction", replacement),
    ]:
        item = doc.add_paragraph()
        item.paragraph_format.left_indent = Inches(0.2)
        item.paragraph_format.space_after = Pt(3)
        lr = item.add_run(f"{label}: ")
        lr.bold = True
        item.add_run(value)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = RGBColor(31, 78, 121)
    return p


def body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("SMART-GROW Defence Paper Compliance Review")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = RGBColor(31, 78, 121)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Exact recommendations, replacement text, sensor parameter tables, and experimental discussion").italic = True
doc.add_paragraph()

heading(doc, "1. Comparison Summary", 1)
body(
    doc,
    "The compliance sheet addresses 29 panel comments, but the current defence paper still needs "
    "technical alignment before the sheet can be considered fully compliant. The largest gaps are "
    "sensor-name consistency, missing environmental parameter tables, unclear controlled vs. uncontrolled "
    "experiment design, and contradictions between the scope and later system/database descriptions.",
)

summary_rows = [
    ["Resolved in paper", "Client/problem discussion, research objectives, locale, respondent discussion, basic instruments, data analysis plan, context/DFD additions, HIPO/use case additions."],
    ["Partially resolved", "ISO 25010 alignment is present but needs objective-to-metric mapping; methodology explains development but not enough experimental protocol detail."],
    ["Not fully resolved", "Compliance row 5 still contains placeholder text; row 27 says sensors were specified, yet paper alternates between SCD40 and DHT22."],
    ["New required additions", "Add sensor parameter tables, ideal temperature/RH thresholds, dashboard visualization table, controlled/uncontrolled environment comparison, and experiment procedure."],
]
table(doc, ["Compliance Status", "Finding"], summary_rows)

heading(doc, "2. Exact Corrections and Replacements", 1)
correction_rows = [
    [
        "Compliance sheet row 5, Actions Taken",
        "Placeholder remains: '------align how?-----'",
        "Replace the whole Actions Taken cell with:",
        "Aligned the IPO output and ISO/IEC 25010 evaluation criteria with the research objectives. Maintainability is measured through maintenance time, fault recovery records, component accessibility, and user feedback on troubleshooting and upkeep.",
    ],
    [
        "Chapter 1, Research Objectives, p. 11",
        "Objective 1 is broad but does not state measurable parameter targets.",
        "Replace Objective 1 with:",
        "To implement an optimized environmental control system integrating automated humidification, ventilation fans, and water refill mechanisms to maintain oyster mushroom fruiting conditions within 85-95% relative humidity and 20-28 deg C temperature, with alerts when readings move outside the acceptable range.",
    ],
    [
        "Chapter 1, Research Objectives, p. 11",
        "Objective 3 says assess current manual operation but does not state comparison basis.",
        "Replace Objective 3 with:",
        "To compare the uncontrolled manual mushroom house and the SMART-GROW-controlled mushroom house in terms of humidity stability, temperature trend, water usage, labor time, system response, and observed mushroom yield/quality indicators.",
    ],
    [
        "Chapter 1, Scope and Limitations, pp. 14-15",
        "Scope says substrate moisture is not monitored, but later DFD/ERD/hardware sections include substrate moisture.",
        "Choose one correction:",
        "If substrate moisture sensor is included: replace 'does not monitor or control light intensity, substrate moisture, or contamination control' with 'does not monitor or control light intensity or contamination; substrate moisture is monitored only as a supporting indicator and is not used as the primary actuator trigger.' If not included: remove substrate moisture from the DFD, ERD, dashboard text, and logs.",
    ],
    [
        "Chapter 3, Research Design, pp. 26-27",
        "Controlled/uncontrolled environments are implied but not named as experiment groups.",
        "Insert after the first paragraph:",
        "The experimental evaluation uses two comparable mushroom-growing areas: an uncontrolled/manual environment and a SMART-GROW-controlled environment. The uncontrolled environment follows ADHIKA's usual manual misting and ventilation practices, while the controlled environment uses automated sensor-based humidification and ventilation. Both areas are monitored using the same logging intervals and evaluation indicators to determine whether SMART-GROW improves environmental stability and operational efficiency.",
    ],
    [
        "Chapter 3, System Architecture / Step 2, p. 71",
        "Paper says DHT22 provides humidity and temperature, but elsewhere says SCD40 handles temperature, humidity, and CO2.",
        "Replace the Step 2 opening sentence with:",
        "The SCD40 sensor module provides continuous carbon dioxide, temperature, and relative humidity readings from the mushroom cultivation area. If a DHT22 sensor is retained, it shall be identified as a secondary temperature-RH reference sensor only, not as the primary environmental sensor.",
    ],
    [
        "Chapter 3, Hardware Requirements, pp. 74-77",
        "Hardware list lacks a consolidated sensor-parameter table.",
        "Insert a new table after Table 3-4:",
        "Add Table 3-5: Sensor Parameters and Ideal Growing Ranges using the parameter table in Section 3 of this document. Renumber the existing software requirements table accordingly.",
    ],
    [
        "Chapter 3, CO2 screen, p. 101",
        "CO2 is displayed as 24%, 437%, and 485%; CO2 should be in ppm for dashboard clarity.",
        "Replace paragraph values with:",
        "The CO2 gauge displays carbon dioxide concentration in parts per million (ppm), such as 437 ppm outside and 485 ppm inside the room. The dashboard uses safe, warning, and critical color bands to guide ventilation decisions.",
    ],
    [
        "Chapter 3, System Testing, p. 105",
        "Testing states target RH but lacks exact pass/fail criteria.",
        "Insert after Functional Testing paragraph:",
        "Functional testing passes when at least 80% of valid readings in the controlled setup remain within 85-95% RH, when temperature remains within the acceptable 20-28 deg C range or is flagged correctly, and when actuator response occurs within the configured response time after a threshold breach.",
    ],
    [
        "Throughout paper",
        "Grammar and terminology inconsistencies: 'will consists', 'firebase', 'useability', 'Figure 3-7shows'.",
        "Apply replacements:",
        "'will consists' -> 'will consist'; 'firebase' -> 'Firebase'; 'useability' -> 'usability'; 'Figure 3-7shows' -> 'Figure 3-7 shows'; 'small-scale mushroom house' -> 'small-scale mushroom houses' or 'a small-scale mushroom house'.",
    ],
]
for idx, row in enumerate(correction_rows, start=1):
    recommendation_block(doc, idx, row[0], row[1], row[2], row[3])

heading(doc, "3. Sensor and Ideal Parameter Tables", 1)
body(
    doc,
    "Use this section as the required visualization/table addition. The RH target follows the paper's current 80-95% basis but is tightened to 85-95% for fruiting discussion, while 80-85% is treated as caution/acceptable rather than ideal. Temperature is presented as an acceptable fruiting range for oyster mushroom cultivation in a warm local setting.",
)
param_rows = [
    ["Relative Humidity", "SCD40 or DHT22 RH channel", "85-95% RH", "80-84% or 96-98% RH", "<80% or >98% RH", "Activate/deactivate humidifier; notify user", "Gauge, line graph, time-in-range bar"],
    ["Temperature", "SCD40 or DHT22 temp channel", "20-28 deg C", "18-19 deg C or 29-30 deg C", "<18 deg C or >30 deg C", "Notify user; activate ventilation/cooling fan if available", "Gauge, trend line, min-max table"],
    ["Carbon Dioxide", "SCD40 CO2 channel", "<1,000 ppm preferred during fruiting", "1,000-1,500 ppm", ">1,500 ppm", "Increase ventilation/fan cycle; notify user", "ppm gauge, color band, event log"],
    ["Water Level", "Water level sensor", "Enough reservoir volume for scheduled misting", "Low level", "Empty / no-flow condition", "Disable pump/humidifier if dry; notify refill", "Tank-level bar and refill alert"],
    ["Substrate Moisture", "Soil/substrate moisture sensor if retained", "Use as supporting indicator only", "Drying trend", "Persistently dry or over-wet reading", "Record observation; do not use as primary actuator unless validated", "Trend table and observation notes"],
]
table(doc, ["Parameter", "Ideal / Warning / Critical", "System Response and Visualization"], [
    [r[0], f"Sensor: {r[1]}\nIdeal: {r[2]}\nWarning: {r[3]}\nCritical: {r[4]}", f"{r[5]}\nVisual: {r[6]}"]
    for r in param_rows
])

spec_rows = [
    ["SCD40", "CO2, temperature, RH", "I2C", "CO2: ppm; Temp: deg C; RH: %", "Use as primary environmental sensor if installed. Keep away from direct mist path to prevent wetting and biased RH readings."],
    ["DHT22", "Temperature, RH", "Single-wire digital", "Temp: deg C; RH: %", "Use only as backup/reference if SCD40 is the main sensor. If DHT22 is removed, delete all DHT22 references."],
    ["Water level sensor", "Reservoir level", "Digital/analog depending on module", "Low/normal or percentage", "Connect to refill and dry-run protection logic."],
    ["Substrate moisture sensor", "Bag/substrate moisture trend", "Analog", "Relative moisture index", "Retain only if the study will discuss substrate readings; otherwise remove from scope, DFD, ERD, and dashboard."],
]
table(doc, ["Sensor", "Technical Use", "Recommendation"], [
    [r[0], f"Measured: {r[1]}\nInterface: {r[2]}\nDisplayed unit: {r[3]}", r[4]]
    for r in spec_rows
])

heading(doc, "4. Controlled vs. Uncontrolled Environment Discussion", 1)
discussion = (
    "Recommended text to insert under Chapter 3, Research Design or Data Gathering Procedure:\n\n"
    "This study compares two cultivation conditions. The uncontrolled environment represents the current ADHIKA practice, where humidity and ventilation are managed manually through grower observation and scheduled spraying. This setup is expected to show wider fluctuations in relative humidity and temperature because intervention depends on grower availability, water pressure, and weather conditions. The controlled environment represents the SMART-GROW setup, where sensor readings are collected continuously and used to activate humidification, ventilation, alerts, and data logging. Comparing these two environments allows the researchers to determine whether automated control improves time-in-range, RH stability, labor efficiency, and consistency of mushroom-growing conditions."
)
body(doc, discussion)

env_rows = [
    ["Control basis", "Manual spraying and grower judgement", "Sensor-based humidifier/fan automation"],
    ["Data captured", "Manual observation log, water use, labor time, yield/quality notes", "Sensor logs, event-response logs, water use, labor time, yield/quality notes"],
    ["Expected pattern", "More frequent RH drops during hot periods and grower absence", "Higher time-in-range and fewer severe RH drops"],
    ["Main comparison metrics", "Baseline for RH stability, labor, and water consumption", "Improvement over baseline using same metrics"],
    ["Risk controls", "Do not alter normal farm practice except measurement logging", "Keep sensor away from direct mist, maintain water level, and document overrides"],
]
table(doc, ["Aspect", "Uncontrolled / Manual Environment", "SMART-GROW Controlled Environment"], env_rows)

heading(doc, "5. Experimental Procedure to Add", 1)
procedure_rows = [
    ["1", "Prepare two comparable growing areas or two equivalent test periods in the same growing house. Label them Manual/Uncontrolled and SMART-GROW/Controlled."],
    ["2", "Place environmental sensors at mushroom canopy height, away from direct water spray and direct fan blast. Use the same sensor placement rule in both setups."],
    ["3", "Record baseline readings for at least 24 hours before automated control begins. Log RH, temperature, CO2 if available, water level, and manual interventions."],
    ["4", "Run the controlled setup using target 85-95% RH and acceptable temperature 20-28 deg C. Log all humidifier/fan activations, alerts, overrides, and connection issues."],
    ["5", "Run the uncontrolled setup using ADHIKA's normal manual misting and ventilation practice. Log spraying time, estimated water used, labor minutes, and observed environmental changes."],
    ["6", "Collect data at fixed intervals, preferably every 5 minutes for sensor readings and daily for water/labor/yield observations."],
    ["7", "Compute time-in-range, RH standard deviation, RH Stability Index, mean temperature, water consumption, labor time, alert count, and response time."],
    ["8", "Compare mushroom indicators such as pinning, cap condition, drying, contamination notes, harvest weight, and grower feedback when available."],
    ["9", "Report limitations such as unequal bag maturity, weather changes, sensor drift, water interruptions, internet outages, or manual overrides."],
]
table(doc, ["Step", "Experimental Action"], procedure_rows)

heading(doc, "6. Compliance Sheet Additions", 1)
sheet_rows = [
    ["Add row after existing row 27", "Specify the actual sensor used for humidity, temperature, and CO2; remove DHT22/SCD40 inconsistency.", "Revised sensor descriptions and system architecture to identify SCD40 as the primary CO2/temperature/RH sensor and DHT22, if retained, as secondary reference only.", "Ch. 3, System Architecture / Hardware Requirements"],
    ["Add row after existing row 27", "Add sensor parameters and ideal growing ranges with visualization/tables.", "Added Sensor Parameters and Ideal Growing Ranges table, including RH, temperature, CO2, water level, actuator response, and dashboard visualization.", "Ch. 3, Hardware Requirements / Data Analysis"],
    ["Add row after existing row 13", "Clarify controlled and uncontrolled environment discussion.", "Added comparison of uncontrolled/manual and SMART-GROW-controlled environments, including expected data sources and comparison metrics.", "Ch. 3, Research Design / Data Gathering Procedure"],
    ["Add row after existing row 18", "Explain experimental procedure for comparing manual and automated setup.", "Added step-by-step experimental protocol, logging intervals, sensor placement, pass/fail criteria, and evaluation metrics.", "Ch. 3, Data Gathering Procedure / System Testing"],
    ["Add row after existing row 24", "Correct CO2 visualization units.", "Changed CO2 dashboard discussion from percentage values to ppm values and added color-band interpretation.", "Ch. 3, Mobile Dashboard / CO2 Monitoring"],
]
table(doc, ["Where to Add", "Comment", "Actions Taken / Page"], [
    [r[0], r[1], f"{r[2]}\nPage/Section: {r[3]}"]
    for r in sheet_rows
])

heading(doc, "7. Suggested Source Notes", 1)
source_rows = [
    ["Paper-internal basis", "The current defence paper already uses 80-95% RH in Chapter 3 testing and Time-in-Range. Keep this as the study's measured range, but state 85-95% RH as the preferred fruiting target."],
    ["Literature support", "Factors affecting mushroom Pleurotus spp. reports Pleurotus growth across about 18-30 deg C and high humidity ranges around 85-95% for cultivation stages."],
    ["Literature support", "Recent oyster mushroom environmental-condition guides commonly list fruiting humidity at 85-95% RH and fruiting temperature around 13-26 deg C depending on strain, with warmer Pleurotus strains tolerating higher values."],
]
table(doc, ["Source Type", "How to Use"], source_rows)

doc.add_paragraph("Web sources checked: https://pmc.ncbi.nlm.nih.gov/articles/PMC6486501/ ; https://www.mdpi.com/2071-1050/17/22/10332")

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
