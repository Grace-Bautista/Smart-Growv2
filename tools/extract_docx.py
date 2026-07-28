import json
import sys
from pathlib import Path

from docx import Document


def cell_text(cell):
    return "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())


def main():
    if len(sys.argv) != 3:
        print("Usage: extract_docx.py <docxPath> <outPath>", file=sys.stderr)
        raise SystemExit(1)

    docx_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2])
    doc = Document(docx_path)

    paragraphs = [
        {"index": i, "text": p.text.strip()}
        for i, p in enumerate(doc.paragraphs)
        if p.text.strip()
    ]
    tables = []
    for ti, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell_text(cell) for cell in row.cells])
        tables.append({"index": ti, "rows": rows})

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(
        json.dumps({"paragraphs": paragraphs, "tables": tables}, indent=2),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
